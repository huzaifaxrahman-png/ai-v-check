"""ATS Resume Checker - Streamlit app powered by Google Gemini.

Upload a resume (PDF or DOCX), optionally paste a job description, and get:
  * an ATS score out of 100 with a category breakdown
  * strengths, problems found, and concrete improvements
  * missing keywords and rewritten bullet suggestions
"""

import io
import json
import os
import re

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

# "gemini-flash-latest" always points to the newest Flash model.
# Override it with a GEMINI_MODEL secret / environment variable if you want.
DEFAULT_MODEL = "gemini-flash-latest"
MAX_RESUME_CHARS = 15000
MAX_JD_CHARS = 6000
MIN_TEXT_CHARS = 100

# Category name -> maximum points (adds up to 100)
CATEGORIES = {
    "keywords": 25,
    "content_quality": 25,
    "formatting": 20,
    "completeness": 20,
    "readability": 10,
}
CATEGORY_LABELS = {
    "keywords": "Keywords & Relevance",
    "content_quality": "Content Quality (impact, metrics)",
    "formatting": "ATS-Friendly Formatting",
    "completeness": "Completeness of Sections",
    "readability": "Readability & Length",
}

SECTION_PATTERNS = {
    "Summary / Objective": r"\b(summary|objective|profile|about me)\b",
    "Experience": r"\b(experience|employment|work history|internship)\b",
    "Education": r"\b(education|academic|qualification)\b",
    "Skills": r"\b(skills|technologies|technical skills|competencies)\b",
    "Projects": r"\b(projects?)\b",
    "Certifications": r"\b(certifications?|certificates?|licenses?)\b",
}


# --------------------------------------------------------------------------
# Text extraction
# --------------------------------------------------------------------------
def extract_text(uploaded_file) -> str:
    """Return plain text from an uploaded PDF or DOCX file."""
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()

    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("This PDF is password protected.")
        pages = [(page.extract_text() or "") for page in reader.pages]
        return "\n".join(pages).strip()

    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text)
        return "\n".join(parts).strip()

    raise ValueError("Unsupported file type. Please upload a PDF or DOCX.")


# --------------------------------------------------------------------------
# Local (rule-based) checks
# --------------------------------------------------------------------------
def basic_checks(text: str) -> dict:
    """Quick deterministic checks that don't need AI."""
    words = re.findall(r"\b\w+\b", text)
    lowered = text.lower()
    return {
        "word_count": len(words),
        "has_email": bool(re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)),
        "has_phone": bool(re.search(r"(\+?\d[\d\s().-]{8,}\d)", text)),
        "has_linkedin": "linkedin.com" in lowered,
        "sections_found": [
            label
            for label, pattern in SECTION_PATTERNS.items()
            if re.search(pattern, lowered)
        ],
    }


# --------------------------------------------------------------------------
# Gemini call
# --------------------------------------------------------------------------
def build_prompt(resume_text: str, job_description: str, checks: dict) -> str:
    jd_block = (
        f"JOB DESCRIPTION:\n{job_description[:MAX_JD_CHARS]}"
        if job_description.strip()
        else "JOB DESCRIPTION: (none provided - judge keywords against common "
        "expectations for the candidate's apparent target role)"
    )
    return f"""You are a strict, fair ATS (Applicant Tracking System) and resume expert.
Evaluate the resume below and return ONLY a JSON object, with no markdown.

Score each category as an integer within its maximum:
- keywords (max 25): relevant hard skills, tools and role keywords; match to the job description if given
- content_quality (max 25): achievements, quantified impact, strong action verbs, clarity
- formatting (max 20): ATS-parsable structure, standard section headings, consistent dates, no signs of tables/columns/graphics breaking parsing
- completeness (max 20): contact info, summary, experience, education, skills present
- readability (max 10): length, concision, grammar, bullet usage

Be realistic: an average resume scores 55-70. Do not inflate scores.

Quick automated checks already run on the resume: {json.dumps(checks)}

Return exactly this JSON shape:
{{
  "scores": {{"keywords": 0, "content_quality": 0, "formatting": 0, "completeness": 0, "readability": 0}},
  "summary": "2-3 sentence overall assessment",
  "strengths": ["..."],
  "issues": ["specific problems found in this resume"],
  "improvements": ["specific, actionable fixes, most important first"],
  "missing_keywords": ["keywords/skills that should be added"],
  "bullet_rewrites": [{{"original": "...", "improved": "..."}}]
}}

Use at most 6 items for strengths, 8 for issues, 8 for improvements, 15 for missing_keywords and 3 for bullet_rewrites.
Base everything on the actual resume text.

{jd_block}

RESUME:
{resume_text[:MAX_RESUME_CHARS]}
"""


def parse_json_response(raw: str) -> dict:
    """Parse model output into a dict, tolerating markdown fences or extra text."""
    cleaned = raw.strip()
    cleaned = re.sub(r"^```(?:json)?\s*|\s*```$", "", cleaned, flags=re.IGNORECASE)
    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        start, end = cleaned.find("{"), cleaned.rfind("}")
        if start != -1 and end > start:
            return json.loads(cleaned[start : end + 1])
        raise


def normalize_result(data: dict) -> dict:
    """Validate and clamp the AI output; compute the total in code."""
    raw_scores = data.get("scores", {}) if isinstance(data, dict) else {}
    scores = {}
    for key, maximum in CATEGORIES.items():
        try:
            value = int(round(float(raw_scores.get(key, 0))))
        except (TypeError, ValueError):
            value = 0
        scores[key] = max(0, min(maximum, value))

    def as_list(key):
        value = data.get(key, [])
        return [str(v) for v in value if str(v).strip()] if isinstance(value, list) else []

    rewrites = []
    for item in data.get("bullet_rewrites", []) if isinstance(data.get("bullet_rewrites"), list) else []:
        if isinstance(item, dict) and item.get("original") and item.get("improved"):
            rewrites.append({"original": str(item["original"]), "improved": str(item["improved"])})

    return {
        "scores": scores,
        "total": sum(scores.values()),
        "summary": str(data.get("summary", "")),
        "strengths": as_list("strengths"),
        "issues": as_list("issues"),
        "improvements": as_list("improvements"),
        "missing_keywords": as_list("missing_keywords"),
        "bullet_rewrites": rewrites,
    }


def analyze_resume(api_key: str, model: str, resume_text: str, job_description: str, checks: dict) -> dict:
    client = genai.Client(api_key=api_key)
    response = client.models.generate_content(
        model=model,
        contents=build_prompt(resume_text, job_description, checks),
        config=types.GenerateContentConfig(
            response_mime_type="application/json",
            temperature=0.2,
        ),
    )
    if not response.text:
        raise RuntimeError("The model returned an empty response. Please try again.")
    return normalize_result(parse_json_response(response.text))


# --------------------------------------------------------------------------
# UI helpers
# --------------------------------------------------------------------------
def get_secret(name: str, default: str = "") -> str:
    """Read from Streamlit secrets first, then environment variables."""
    try:
        if name in st.secrets:
            return str(st.secrets[name])
    except Exception:
        pass  # no secrets file present
    return os.environ.get(name, default)


def score_label(score: int) -> str:
    if score >= 80:
        return "Excellent - highly ATS-friendly"
    if score >= 65:
        return "Good - a few fixes will help"
    if score >= 50:
        return "Fair - needs noticeable improvement"
    return "Poor - major improvements needed"


def render_list(title: str, items: list):
    if not items:
        return
    st.subheader(title)
    for item in items:
        st.markdown(f"- {item}")


def render_results(result: dict, checks: dict):
    total = result["total"]
    st.divider()
    col1, col2 = st.columns([1, 2])
    with col1:
        st.metric("ATS Score", f"{total} / 100")
        st.caption(score_label(total))
    with col2:
        st.progress(total / 100)
        if result["summary"]:
            st.write(result["summary"])

    st.subheader("Score breakdown")
    for key, maximum in CATEGORIES.items():
        value = result["scores"][key]
        st.write(f"**{CATEGORY_LABELS[key]}**: {value} / {maximum}")
        st.progress(value / maximum)

    render_list("Strengths", result["strengths"])
    render_list("Problems found", result["issues"])
    render_list("How to improve", result["improvements"])

    if result["missing_keywords"]:
        st.subheader("Missing keywords")
        st.write(", ".join(f"`{k}`" for k in result["missing_keywords"]))

    if result["bullet_rewrites"]:
        st.subheader("Suggested bullet rewrites")
        for item in result["bullet_rewrites"]:
            st.markdown(f"**Before:** {item['original']}")
            st.markdown(f"**After:** {item['improved']}")
            st.write("")

    with st.expander("Automated checks"):
        st.write(f"Word count: **{checks['word_count']}**")
        st.write(f"Email found: **{'Yes' if checks['has_email'] else 'No'}**")
        st.write(f"Phone found: **{'Yes' if checks['has_phone'] else 'No'}**")
        st.write(f"LinkedIn found: **{'Yes' if checks['has_linkedin'] else 'No'}**")
        st.write("Sections detected: " + (", ".join(checks["sections_found"]) or "none"))


# --------------------------------------------------------------------------
# Main app
# --------------------------------------------------------------------------
def main():
    st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="centered")
    st.title("📄 ATS Resume Checker")
    st.write("Upload your resume to get an ATS score and tips to improve it.")

    api_key = get_secret("GEMINI_API_KEY")
    model = get_secret("GEMINI_MODEL", DEFAULT_MODEL)

    with st.sidebar:
        st.header("Settings")
        if not api_key:
            api_key = st.text_input("Gemini API key", type="password",
                                    help="Get a free key at https://aistudio.google.com/apikey")
        else:
            st.success("API key loaded")
        model = st.text_input("Model", value=model)
        st.caption("Your resume is sent to Google's Gemini API for analysis.")

    uploaded = st.file_uploader("Upload resume (PDF or DOCX)", type=["pdf", "docx"])
    job_description = st.text_area(
        "Job description (optional, for a better keyword match)",
        height=150,
        placeholder="Paste the job posting here...",
    )

    if st.button("Analyze resume", type="primary"):
        if not api_key:
            st.error("Please provide a Gemini API key in the sidebar.")
            return
        if uploaded is None:
            st.error("Please upload a resume first.")
            return

        try:
            with st.spinner("Reading your resume..."):
                text = extract_text(uploaded)
        except Exception as exc:
            st.error(f"Could not read the file: {exc}")
            return

        if len(text) < MIN_TEXT_CHARS:
            st.error(
                "Almost no text could be extracted. If your resume is a scanned image, "
                "that is itself an ATS problem - export a text-based PDF or DOCX instead."
            )
            return

        checks = basic_checks(text)
        try:
            with st.spinner("Analyzing with Gemini..."):
                result = analyze_resume(api_key, model.strip() or DEFAULT_MODEL,
                                        text, job_description, checks)
        except json.JSONDecodeError:
            st.error("The AI returned an unreadable response. Please try again.")
            return
        except Exception as exc:
            st.error(f"Analysis failed: {exc}")
            return

        render_results(result, checks)


if __name__ == "__main__":
    main()
